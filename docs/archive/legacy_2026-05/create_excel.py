import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# 创建工作簿
wb = openpyxl.Workbook()

# ===== 3-shot 结果表 =====
ws3 = wb.active
ws3.title = "3-shot Results"

# 表头
headers = ["Model", "Class", "P", "R", "mAP50", "mAP50-95"]
for col, header in enumerate(headers, 1):
    cell = ws3.cell(row=1, column=col, value=header)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center')

# 3-shot 数据
data_3shot = [
    ["novel_finetune", "all", 0.178, 0.414, 0.205, 0.154],
    ["", "bird", 0.138, 0.19, 0.11, 0.0696],
    ["", "bus", 0.218, 0.653, 0.493, 0.415],
    ["", "cow", 0.179, 0.43, 0.134, 0.0913],
    ["", "motorbike", 0.194, 0.388, 0.146, 0.0939],
    ["", "sofa", 0.162, 0.408, 0.142, 0.101],
    ["novel_finetune_cosine", "all", 0.0168, 0.838, 0.513, 0.349],
    ["", "bird", 0.0111, 0.704, 0.374, 0.227],
    ["", "bus", 0.0313, 0.901, 0.765, 0.629],
    ["", "cow", 0.00162, 0.811, 0.43, 0.256],
    ["", "motorbike", 0.00819, 0.8, 0.552, 0.345],
    ["", "sofa", 0.0318, 0.975, 0.443, 0.288],
    ["novel_finetune_cosine_proto", "all", 0.0237, 0.834, 0.523, 0.339],
    ["", "bird", 0.0259, 0.68, 0.382, 0.212],
    ["", "bus", 0.0266, 0.883, 0.739, 0.579],
    ["", "cow", 0.0162, 0.848, 0.544, 0.338],
    ["", "motorbike", 0.0274, 0.803, 0.423, 0.235],
    ["", "sofa", 0.0226, 0.954, 0.525, 0.331],
    ["novel_finetune_cosine_fused", "all", 0.014, 0.839, 0.539, 0.352],
    ["", "bird", 0.0179, 0.702, 0.4, 0.223],
    ["", "bus", 0.015, 0.873, 0.76, 0.586],
    ["", "cow", 0.0087, 0.836, 0.537, 0.342],
    ["", "motorbike", 0.0135, 0.812, 0.472, 0.278],
    ["", "sofa", 0.0149, 0.971, 0.525, 0.331],
]

for row_idx, row_data in enumerate(data_3shot, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws3.cell(row=row_idx, column=col_idx, value=value)
        if col_idx >= 3:  # 数值列右对齐
            cell.alignment = Alignment(horizontal='right')
            cell.number_format = '0.0000'

# 调整列宽
ws3.column_dimensions['A'].width = 30
ws3.column_dimensions['B'].width = 12
for col in ['C', 'D', 'E', 'F']:
    ws3.column_dimensions[col].width = 12

# ===== 5-shot 结果表 =====
ws5 = wb.create_sheet("5-shot Results")

for col, header in enumerate(headers, 1):
    cell = ws5.cell(row=1, column=col, value=header)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center')

data_5shot = [
    ["novel_finetune", "all", 0.261, 0.434, 0.329, 0.242],
    ["", "bird", 0.197, 0.137, 0.15, 0.0958],
    ["", "bus", 0.305, 0.775, 0.7, 0.579],
    ["", "cow", 0.405, 0.484, 0.393, 0.261],
    ["", "motorbike", 0.252, 0.474, 0.264, 0.17],
    ["", "sofa", 0.144, 0.301, 0.14, 0.106],
    ["novel_finetune_cosine", "all", 0.0124, 0.907, 0.669, 0.455],
    ["", "bird", 0.00826, 0.752, 0.425, 0.238],
    ["", "bus", 0.0226, 0.962, 0.837, 0.661],
    ["", "cow", 0.00201, 0.943, 0.725, 0.452],
    ["", "motorbike", 0.00815, 0.892, 0.714, 0.459],
    ["", "sofa", 0.0208, 0.987, 0.646, 0.466],
    ["novel_finetune_cosine_proto", "all", 0.0534, 0.866, 0.688, 0.469],
    ["", "bird", 0.0552, 0.712, 0.502, 0.282],
    ["", "bus", 0.0568, 0.92, 0.815, 0.65],
    ["", "cow", 0.0617, 0.914, 0.769, 0.505],
    ["", "motorbike", 0.0488, 0.818, 0.664, 0.421],
    ["", "sofa", 0.0447, 0.967, 0.688, 0.488],
    ["novel_finetune_cosine_fused", "all", 0.0636, 0.868, 0.69, 0.471],
    ["", "bird", 0.0602, 0.717, 0.488, 0.28],
    ["", "bus", 0.0763, 0.911, 0.813, 0.649],
    ["", "cow", 0.0598, 0.91, 0.763, 0.498],
    ["", "motorbike", 0.0633, 0.837, 0.691, 0.436],
    ["", "sofa", 0.0586, 0.967, 0.694, 0.492],
]

for row_idx, row_data in enumerate(data_5shot, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws5.cell(row=row_idx, column=col_idx, value=value)
        if col_idx >= 3:
            cell.alignment = Alignment(horizontal='right')
            cell.number_format = '0.0000'

ws5.column_dimensions['A'].width = 30
ws5.column_dimensions['B'].width = 12
for col in ['C', 'D', 'E', 'F']:
    ws5.column_dimensions[col].width = 12

# ===== 对比汇总表 =====
ws_compare = wb.create_sheet("Comparison")

compare_headers = ["Model", "Shot", "P", "R", "mAP50", "mAP50-95"]
for col, header in enumerate(compare_headers, 1):
    cell = ws_compare.cell(row=1, column=col, value=header)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal='center')

compare_data = [
    ["novel_finetune", "3-shot", 0.178, 0.414, 0.205, 0.154],
    ["novel_finetune", "5-shot", 0.261, 0.434, 0.329, 0.242],
    ["novel_finetune_cosine", "3-shot", 0.0168, 0.838, 0.513, 0.349],
    ["novel_finetune_cosine", "5-shot", 0.0124, 0.907, 0.669, 0.455],
    ["novel_finetune_cosine_proto", "3-shot", 0.0237, 0.834, 0.523, 0.339],
    ["novel_finetune_cosine_proto", "5-shot", 0.0534, 0.866, 0.688, 0.469],
    ["novel_finetune_cosine_fused", "3-shot", 0.014, 0.839, 0.539, 0.352],
    ["novel_finetune_cosine_fused", "5-shot", 0.0636, 0.868, 0.69, 0.471],
]

for row_idx, row_data in enumerate(compare_data, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws_compare.cell(row=row_idx, column=col_idx, value=value)
        if col_idx >= 3:
            cell.alignment = Alignment(horizontal='right')
            cell.number_format = '0.0000'

ws_compare.column_dimensions['A'].width = 30
ws_compare.column_dimensions['B'].width = 10
for col in ['C', 'D', 'E', 'F']:
    ws_compare.column_dimensions[col].width = 12

# 保存文件
output_path = "e:/FSOD_Baseline/FSOD_LLM/FSOD_results.xlsx"
wb.save(output_path)
print(f"Excel文件已保存到: {output_path}")
